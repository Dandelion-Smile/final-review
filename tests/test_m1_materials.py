from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from PIL import Image

from final_review.agent import citation_issues
from final_review.api import create_app
from final_review.material_conversion import convert_upload
from final_review.material_jobs import process_material_job
from final_review.policy import SOURCE_PRIORITY
from final_review.schemas import SourceType


def _image(format_name: str) -> bytes:
    output = BytesIO()
    Image.new("RGB", (80, 40), "white").save(output, format=format_name)
    return output.getvalue()


def _run_queued_job(system):
    job = system.store.claim_material_job()
    assert job is not None
    process_material_job(system.store, system.kb, job,
                         system.settings.max_upload_mb * 1024 * 1024)
    return system.store.get_material_job(job["job_id"])


def test_upload_survives_worker_with_different_working_directory(system, tmp_path, monkeypatch):
    system.settings.uploads_dir = str(tmp_path / "uploads")
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "课程笔记", "source_type": "homework"},
            files={"file": ("notes.md", "网络协议与三次握手".encode())},
        )
        assert response.status_code == 202
        document = system.store.get("document", response.json()["document_id"])
        assert Path(document["file_path"]).is_absolute()
        monkeypatch.chdir(tmp_path.parent)
        assert _run_queued_job(system)["status"] == "succeeded"


@pytest.mark.parametrize(
    ("suffix", "format_name"),
    [("png", "PNG"), ("jpg", "JPEG"), ("jpeg", "JPEG"), ("webp", "WEBP")],
)
def test_image_ocr_upload_is_searchable(system, tmp_path, monkeypatch, suffix, format_name):
    from final_review import material_conversion

    system.settings.uploads_dir = str(tmp_path / "uploads")
    monkeypatch.setattr(material_conversion.shutil, "which", lambda name: "tesseract")

    def run_ocr(*args, **kwargs):
        assert kwargs["encoding"] == "utf-8"
        return SimpleNamespace(returncode=0, stdout="学 习 通 作 业：网络三次握手")

    monkeypatch.setattr(
        material_conversion.subprocess,
        "run",
        run_ocr,
    )
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "学习通截图", "source_type": "homework"},
            files={"file": (f"homework.{suffix}", _image(format_name))},
        )
        assert response.status_code == 202, response.text
        assert _run_queued_job(system)["status"] == "succeeded"
        document = system.store.get("document", response.json()["document_id"])
        assert document["parse_status"] == "ready"
        assert document["source_origin"] == "user_upload"
        assert "学习通作业" in document["cleaned_markdown"]
        assert any(chunk["document_id"] == document["document_id"] for chunk in system.store.chunks)


def test_bad_image_has_explainable_failure_and_no_chunks(system, tmp_path):
    system.settings.uploads_dir = str(tmp_path / "uploads")
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "损坏截图", "source_type": "other_practice"},
            files={"file": ("broken.png", b"not a png")},
        )
        assert response.status_code == 202
        job = _run_queued_job(system)
        assert job["status"] == "failed"
        assert "图片" in job["error_message"]
        document = client.get("/api/courses/net/documents").json()["items"][-1]
        assert document["parse_status"] == "failed"
        assert document["parse_error"] == job["error_message"]
        assert not any(
            chunk["document_id"] == document["document_id"] for chunk in system.store.chunks
        )


def test_picture_only_ppt_has_explainable_failure_and_no_chunks(system, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    system.settings.uploads_dir = str(tmp_path / "uploads")
    picture = BytesIO(_image("PNG"))
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_picture(picture, Inches(1), Inches(1))
    ppt = BytesIO()
    slides.save(ppt)
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "图片课件", "source_type": "teacher_ppt"},
            files={"file": ("pictures.pptx", ppt.getvalue())},
        )
        assert response.status_code == 202
        job = _run_queued_job(system)
        assert job["status"] == "failed"
        assert "图片单独上传" in job["error_message"]
        document_id = response.json()["document_id"]
        assert client.get(f"/api/courses/net/documents/{document_id}/chunks").status_code == 404
        assert not any(chunk["document_id"] == document_id for chunk in system.store.chunks)


@pytest.mark.parametrize(
    "suffix",
    ["md", "txt", "pdf", "ppt", "pptx", "doc", "docx", "png", "jpg", "jpeg", "webp"],
)
def test_each_supported_format_has_explainable_failure(system, tmp_path, monkeypatch, suffix):
    from final_review import material_conversion

    system.settings.uploads_dir = str(tmp_path / "uploads")
    monkeypatch.setattr(material_conversion, "_find_executable", lambda *_args: None)
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "损坏资料", "source_type": "other_practice"},
            files={"file": (f"broken.{suffix}", b"\xffbroken")},
        )
        assert response.status_code == 202
        job = _run_queued_job(system)
        assert job["status"] == "failed"
        assert job["error_message"]
        document = client.get("/api/courses/net/documents").json()["items"][-1]
        assert document["parse_status"] == "failed"
        assert document["parse_error"] == job["error_message"]


def test_empty_ocr_result_stays_failed(system, tmp_path, monkeypatch):
    from final_review import material_conversion

    system.settings.uploads_dir = str(tmp_path / "uploads")
    monkeypatch.setattr(material_conversion.shutil, "which", lambda name: "tesseract")
    monkeypatch.setattr(
        material_conversion.subprocess,
        "run",
        lambda *args, **kwargs: SimpleNamespace(returncode=0, stdout="  "),
    )
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "空白截图", "source_type": "homework"},
            files={"file": ("blank.png", _image("PNG"))},
        )
        assert response.status_code == 202
        job = _run_queued_job(system)
        assert job["status"] == "failed"
        assert "没有识别到" in job["error_message"]
        document = client.get("/api/courses/net/documents").json()["items"][-1]
        assert document["parse_status"] == "failed"
        assert not any(
            chunk["document_id"] == document["document_id"] for chunk in system.store.chunks
        )


def test_index_failure_does_not_leave_parsing_or_searchable_chunks(system, tmp_path, monkeypatch):
    system.settings.uploads_dir = str(tmp_path / "uploads")

    def fail(_texts):
        raise RuntimeError("embedding unavailable")

    monkeypatch.setattr(system.kb.embeddings, "embed_documents", fail)
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "索引故障", "source_type": "homework"},
            files={"file": ("note.md", "网络协议".encode())},
        )
        assert response.status_code == 202
        for _ in range(3):
            job = _run_queued_job(system)
        assert job["status"] == "failed"
        document = client.get("/api/courses/net/documents").json()["items"][-1]
        assert document["parse_status"] == "failed"
        assert "暂时不可用" in document["parse_error"]
        assert not any(
            chunk["document_id"] == document["document_id"] for chunk in system.store.chunks
        )


@pytest.mark.parametrize(("suffix", "converted_suffix"), [("doc", "docx"), ("ppt", "pptx")])
def test_legacy_office_conversion_uses_bounded_local_tool(
    monkeypatch, tmp_path, suffix, converted_suffix
):
    from final_review import material_conversion

    if suffix == "doc":
        from zipfile import ZipFile

        converted = tmp_path / "sample.docx"
        with ZipFile(converted, "w") as archive:
            archive.writestr("word/document.xml", "<document>网络三次握手</document>")
        monkeypatch.setattr(
            material_conversion,
            "MarkItDown",
            lambda **kwargs: SimpleNamespace(
                convert=lambda path: SimpleNamespace(text_content="网络三次握手")
            ),
        )
    else:
        from pptx import Presentation

        office = Presentation()
        slide = office.slides.add_slide(office.slide_layouts[1])
        slide.shapes.title.text = "网络三次握手"
        converted = tmp_path / "sample.pptx"
        office.save(converted)
    monkeypatch.setattr(material_conversion.shutil, "which", lambda name: "soffice")

    def run(command, **kwargs):
        assert "--headless" in command
        assert kwargs["timeout"] == material_conversion.LEGACY_OFFICE_TIMEOUT_SECONDS
        Path(command[-1]).with_suffix(f".{converted_suffix}").write_bytes(converted.read_bytes())
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(material_conversion.subprocess, "run", run)
    markdown = convert_upload(b"legacy office bytes", f"sample.{suffix}", 10_000_000)
    assert "网络三次握手" in markdown


def test_source_categories_keep_existing_priority():
    assert SOURCE_PRIORITY[SourceType.homework] > SOURCE_PRIORITY[SourceType.other_practice]
    assert SOURCE_PRIORITY[SourceType.other_practice] == SOURCE_PRIORITY[SourceType.crash_course]
    assert SOURCE_PRIORITY[SourceType.crash_course] > SOURCE_PRIORITY[SourceType.ai_supplement]


def test_ai_supplement_keeps_its_label(system, tmp_path):
    system.settings.uploads_dir = str(tmp_path / "uploads")
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "AI 整理", "source_type": "ai_supplement"},
            files={"file": ("note.txt", "AI 整理内容".encode())},
        )
        assert response.status_code == 202, response.text
        assert _run_queued_job(system)["status"] == "succeeded"
        document = system.store.get("document", response.json()["document_id"])
        assert document["source_type"] == "ai_supplement"
        assert document["source_origin"] == "user_upload"

    issues = citation_issues(
        {"source_type": "teacher_ppt", "citations": [{"chunk_id": "ai", "quote": "补充"}]},
        [{"chunk_id": "ai", "content": "AI 补充", "source_type": "ai_supplement"}],
    )
    assert "题源标记不匹配引用" in issues
