from fastapi.testclient import TestClient

from final_review.api import create_app
from final_review.config import ChatModelConfig
from final_review.material_jobs import process_material_job


def _run_queued_job(system):
    job = system.store.claim_material_job()
    assert job is not None
    process_material_job(system.store, system.kb, job,
                         system.settings.max_upload_mb * 1024 * 1024)
    return system.store.get_material_job(job["job_id"])


def test_api_full_feedback_cycle(system):
    with TestClient(create_app(system.settings, system)) as client:
        assert client.get("/health").status_code == 200
        assert client.get("/openapi.json").status_code == 200
        response = client.post(
            "/agent/invoke",
            json={
                "course_id": "net",
                "session_id": "http",
                "message": "出题",
                "intent": "quiz",
            },
        )
        assert response.json()["status"] == "needs_input"
        result = client.post(
            "/agent/resume",
            json={
                "course_id": "net",
                "session_id": "http",
                "exam_profile": {"question_types": ["short_answer"]},
            },
        ).json()
        assert result["status"] == "awaiting_answers"
        assert "reference_answer" not in str(result)
        graded = client.post(
            "/assessment/evaluate",
            json={
                "course_id": "net",
                "session_id": "http",
                "answers": {q["id"]: "不清楚" for q in result["questions"]},
            },
        )
        assert graded.status_code == 200
        assert graded.json()["assessment"]["score"] == 30


def test_upload_supported_and_reject_bad_file(system, tmp_path):
    system.settings.uploads_dir = str(tmp_path / "uploads")
    with TestClient(create_app(system.settings, system)) as client:
        data = {"course_id": "net", "title": "上传", "source_type": "homework"}
        assert (
            client.post(
                "/knowledge/upload", data=data, files={"file": ("note.md", "测试内容".encode())}
            ).status_code
            == 202
        )
        assert (
            client.post(
                "/knowledge/upload", data=data, files={"file": ("test.exe", b"MZ")}
            ).status_code
            == 422
        )
        assert _run_queued_job(system)["status"] == "succeeded"
        assert (
            client.post(
                "/knowledge/upload", data=data, files={"file": ("empty.md", b"")}
            ).status_code
            == 422
        )


def test_auth_and_missing_session(system):
    from pydantic import SecretStr

    system.settings.api_token = SecretStr("test-token")
    with TestClient(create_app(system.settings, system)) as client:
        url = "/agent/recover?course_id=net&session_id=missing"
        assert client.post(url).status_code == 401
        assert client.post(url, headers={"Authorization": "Bearer test-token"}).status_code == 404


def test_invalid_request_never_reaches_model(system):
    with TestClient(create_app(system.settings, system)) as client:
        result = client.post(
            "/agent/invoke",
            json={
                "course_id": "net",
                "session_id": "test",
                "message": "出题",
                "intent": "unknown",
            },
        )
        assert result.status_code == 422
        assert system.model.retrieval_calls == 0


def test_chat_uses_knowledge_service_and_persists_messages(system):
    system.settings.chat_models = [
        ChatModelConfig(
            id="review",
            label="复习模型",
            model="test-model",
            base_url="https://example.invalid/v1",
            api_key="test-key",
        )
    ]
    with TestClient(create_app(system.settings, system)) as client:
        models = client.get("/api/chat/models")
        assert models.json() == {"items": [{"id": "review", "label": "复习模型"}]}
        response = client.post("/api/chat", json={"message": "帮我复习需求分析"})
        assert response.status_code == 200
        assert response.json()["reply"]
        # The phase-two API intentionally rejects the old unscoped route.
        assert client.get("/api/conversations/default/messages").status_code == 404
        messages = client.get(
            "/api/courses/software-engineering-basics/conversations/default/messages"
        )
        assert messages.status_code == 200
        assert {item["role"] for item in messages.json()["items"]} == {"user", "assistant"}
        assert (
            client.post(
                "/api/chat", json={"message": "测试", "model_id": "not-allowed"}
            ).status_code
            == 422
        )


def test_chat_greeting_does_not_require_course_evidence(system):
    from final_review.config import ChatModelConfig

    system.settings.chat_models = [
        ChatModelConfig(
            id="review",
            label="review",
            model="test-model",
            base_url="https://example.invalid/v1",
            api_key="test-key",
        )
    ]
    with TestClient(create_app(system.settings, system)) as client:
        reply = client.post("/api/chat", json={"message": "你好"}).json()["reply"]
        assert "准备好了" in reply


def test_direct_chat_uses_selected_model_and_history(system, monkeypatch):
    system.settings.chat_models = [
        ChatModelConfig(
            id="deepseek", label="DeepSeek", model="deepseek-test",
            base_url="https://api.deepseek.com", api_key="test-key",
        ),
        ChatModelConfig(
            id="gemini", label="Gemini", model="gemini-test",
            base_url="https://example.invalid/v1", api_key="test-key", grounded=False,
        ),
    ]
    calls = []

    class FakeCompletions:
        def create(self, **kwargs):
            calls.append(kwargs)
            from types import SimpleNamespace
            return SimpleNamespace(
                choices=[SimpleNamespace(message=SimpleNamespace(content="真实回复"))]
            )

    class FakeOpenAI:
        def __init__(self, **kwargs):
            self.chat = type("Chat", (), {"completions": FakeCompletions()})()

    monkeypatch.setattr("final_review.api.OpenAI", FakeOpenAI)
    with TestClient(create_app(system.settings, system)) as client:
        result = client.post("/api/chat", json={
            "message": "你好", "model_id": "gemini", "mode": "direct",
            "history": [{"role": "user", "content": "先前的问题"},
                        {"role": "assistant", "content": "先前的回答"}],
        })
        assert result.status_code == 200
        assert result.json()["reply"] == "真实回复"
        assert result.json()["model"] == "Gemini"
        assert calls[0]["model"] == "gemini-test"
        assert [item["role"] for item in calls[0]["messages"]] == ["system", "user"]
        continued = client.post("/api/chat", json={
            "message": "继续", "model_id": "gemini", "mode": "direct",
            "history": [{"role": "user", "content": "不可信的页面历史"}],
        })
        assert continued.status_code == 200
        assert [item["role"] for item in calls[1]["messages"]] == [
            "system", "user", "assistant", "user"
        ]
        assert calls[1]["messages"][1]["content"] == "你好"
        deepseek_reply = client.post("/api/chat", json={
            "message": "继续", "model_id": "deepseek", "mode": "direct",
        })
        assert deepseek_reply.status_code == 200
        assert calls[2]["extra_body"] == {"thinking": {"type": "disabled"}}
        assert "extra_body" not in calls[1]
        conversations = client.get(
            "/api/courses/software-engineering-basics/conversations"
        ).json()["items"]
        assert conversations[0]["title"] == "你好"
        renamed = client.patch(
            "/api/courses/software-engineering-basics/conversations/default",
            json={"title": "  期末复习重点  "},
        )
        assert renamed.status_code == 200
        assert renamed.json()["title"] == "期末复习重点"
        assert client.get(
            "/api/courses/software-engineering-basics/conversations"
        ).json()["items"][0]["title"] == "期末复习重点"
        assert client.patch(
            "/api/courses/software-engineering-basics/conversations/default",
            json={"title": "   "},
        ).status_code == 422
        other_course = client.post("/api/courses", json={"name": "Other"}).json()["course_id"]
        assert (
            client.get(f"/api/courses/{other_course}/conversations/default/messages").status_code
            == 404
        )
        assert client.patch(
            f"/api/courses/{other_course}/conversations/default", json={"title": "不可改"}
        ).status_code == 404


def test_real_pptx_conversion(system, tmp_path):
    from io import BytesIO

    from pptx import Presentation

    system.settings.uploads_dir = str(tmp_path / "uploads")

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "TCP 三次握手"
    slide.placeholders[1].text = "同步双方初始序列号并确认双方收发能力。"
    data = BytesIO()
    slides.save(data)
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post(
            "/knowledge/upload",
            data={
                "course_id": "net",
                "title": "转换测试",
                "source_type": "teacher_ppt",
            },
            files={"file": ("slides.pptx", data.getvalue())},
        )
        assert response.status_code == 202, response.text
        assert _run_queued_job(system)["status"] == "succeeded"
        stored = system.store.get("document", response.json()["document_id"])
        assert "三次握手" in stored["cleaned_markdown"]


def test_phase_one_persistence_endpoints(system, tmp_path):
    """Courses, local source files, scoped chat history and reports form the local MVP."""
    from final_review.config import ChatModelConfig

    system.settings.uploads_dir = str(tmp_path / "uploads")
    system.settings.chat_models = [
        ChatModelConfig(
            id="review",
            label="review",
            model="test-model",
            base_url="https://example.invalid/v1",
            api_key="test-key",
        )
    ]
    with TestClient(create_app(system.settings, system)) as client:
        course = client.post("/api/courses", json={"name": "Networks"}).json()
        assert client.get("/api/courses").json()["items"] == [course]
        uploaded = client.post(
            "/knowledge/upload",
            data={"course_id": course["course_id"], "title": "note", "source_type": "homework"},
            files={"file": ("note.md", "TCP 三次握手".encode())},
        ).json()
        assert _run_queued_job(system)["status"] == "succeeded"
        document = client.get(f"/api/courses/{course['course_id']}/documents").json()["items"][0]
        assert document["parse_status"] == "ready"
        assert (tmp_path / "uploads" / "local-user" / course["course_id"]).exists()
        assert (
            client.post(
                "/api/chat",
                json={
                    "course_id": course["course_id"],
                    "conversation_id": "first",
                    "message": "复习",
                },
            ).status_code
            == 200
        )
        history = client.get(
            f"/api/courses/{course['course_id']}/conversations/first/messages"
        ).json()["items"]
        assert [row["role"] for row in history] == ["user", "assistant"]
        preview = client.post(
            f"/api/courses/{course['course_id']}/documents/{uploaded['document_id']}/deletion-preview"
        ).json()
        assert client.post(
            f"/api/courses/{course['course_id']}/documents/{uploaded['document_id']}/delete",
            json={"confirmation_id": preview["confirmation_id"]},
        ).json() == {"deleted": True, "retained_source_snapshot": False}


def test_fast_quiz_uses_direct_retrieval_and_persists_attempt(system, monkeypatch):
    from final_review.config import ChatModelConfig

    system.settings.chat_models = [
        ChatModelConfig(
            id="review",
            label="review",
            model="test-model",
            base_url="https://example.invalid/v1",
            api_key="test-key",
        )
    ]

    class FastModel:
        def fast_quiz(self, data):
            return {
                "questions": [
                    {
                        "id": "model-id",
                        "knowledge_point": "三次握手",
                        "question_type": "short_answer",
                        "stem": "为什么需要三次握手？",
                        "options": [],
                        "reference_answer": "同步序列号",
                        "explanation": "双方确认收发能力",
                        "must_include": ["序列号"],
                        "source_chunk_ids": [data["evidence"][0]["chunk_id"]],
                    }
                ]
            }

        def grade(self, data):
            question_id = data["quiz"]["questions"][0]["id"]
            return {
                "items": [
                    {
                        "question_id": question_id,
                        "score": 80,
                        "error_type": "",
                        "feedback": "正确",
                        "missing_points": [],
                    }
                ]
            }

    monkeypatch.setattr("final_review.api.build_fast_quiz_model", lambda *_: FastModel())
    with TestClient(create_app(system.settings, system)) as client:
        generated = client.post(
            "/api/quiz/generate",
            json={
                "course_id": "net",
                "chapter": "TCP",
                "question_types": ["short_answer"],
                "question_count": 1,
                "model_id": "review",
            },
        ).json()
        question = generated["questions"][0]
        graded = client.post(
            f"/api/quiz/{generated['session_id']}/submit",
            json={"course_id": "net", "answers": {question["id"]: "同步序列号"}},
        )
        assert graded.status_code == 200
        assert graded.json()["assessment"]["score"] == 80
        assert client.get("/api/courses/net/attempts").json()["items"]
        detail = client.get(f"/api/courses/net/attempts/{generated['session_id']}").json()
        assert detail["answers"] == {question["id"]: "同步序列号"}
        repeated = client.post(
            f"/api/quiz/{generated['session_id']}/submit",
            json={"course_id": "net", "answers": {question["id"]: "不同答案"}},
        ).json()
        assert repeated["already_graded"] is True
