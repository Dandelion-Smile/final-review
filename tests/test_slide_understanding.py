from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from final_review.api import create_app
from final_review.material_conversion import ConvertedMaterial
from final_review.slide_understanding import SlideInterpreter


def interpreter(replies):
    instance = SlideInterpreter.__new__(SlideInterpreter)
    instance.identity = "test-model-and-prompt"
    instance.model = "test-vision"
    instance.repairs = 1
    instance.calls = []

    def ask(prompt, image, text):
        instance.calls.append((prompt, image, text))
        return next(replies)

    instance._ask = ask
    return instance


def reading(**changes):
    return {
        "kind": "knowledge",
        "title": "请求处理",
        "blocks": [{"title": "请求对象", "text": "容器创建请求对象，并调用 service() 方法。"}],
        "uncertainties": [],
        **changes,
    }


def review(**changes):
    return {"faithful": True, "complete": True, "readable": True, "issues": [], **changes}


def test_verified_page_cache_avoids_repeated_model_calls(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image bytes")
    reader = interpreter(iter([reading(), review()]))
    result = reader.read(image, "原生文字", "备注", tmp_path / "cache")
    assert result["quality"] == "verified"
    assert len(reader.calls) == 2
    assert reader.calls[0][1] == b"image bytes"
    assert reader.read(image, "原生文字", "备注", tmp_path / "cache") == result
    assert len(reader.calls) == 2
    reader.identity = "changed-model"
    with pytest.raises(StopIteration):
        reader.read(image, "原生文字", "备注", tmp_path / "cache")


def test_unfaithful_page_repairs_then_remains_excluded(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    reader = interpreter(
        iter(
            [
                reading(),
                review(faithful=False, issues=["图中没有此步骤"]),
                reading(uncertainties=["箭头不清楚"]),
                review(),
            ]
        )
    )
    result = reader.read(image, "文字", "", tmp_path / "cache")
    assert result["quality"] == "review_needed"
    assert result["issues"] == ["箭头不清楚"]
    assert "图中没有此步骤" in reader.calls[2][2]


def test_corrupt_model_output_is_not_accepted_even_if_reviewer_passes(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    broken = reading(blocks=[{"title": "流程", "markdown": "乱码\ufffd service()()"}])
    reader = interpreter(iter([broken, review(), broken, review()]))
    assert reader.read(image, "", "", tmp_path / "cache")["quality"] == "review_needed"


def test_navigation_page_is_verified_without_knowledge_blocks(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    reader = interpreter(iter([reading(kind="navigation", blocks=[]), review()]))
    result = reader.read(image, "目录", "", tmp_path / "cache")
    assert result["quality"] == "verified"
    assert result["blocks"] == []


def test_final_plain_text_keeps_code_symbols_and_removes_prose_markup():
    from final_review.plain_material_text import plain_material_text

    source = "# 标题\n\n**定义**与 `service()`\n```css\n#main { color: #fff; }\n```\n- 普通条目"
    assert (
        plain_material_text(source) == "标题\n\n定义与 service()\n#main { color: #fff; }\n普通条目"
    )


def test_understood_plain_code_is_not_stripped_again(system):
    from final_review.schemas import MaterialInput

    source = "示例代码\n# Python comment\nx = 2 ** 3\n#main { color: #fff; }"
    document, chunks = system.kb.prepare(MaterialInput(
        course_id="net", title="代码页", source_type="teacher_ppt", markdown=source,
    ), sections=[{"text": source, "position_kind": "slide", "position": 1}], plain_text=True)
    assert chunks[0]["content"] == source
    assert document["cleaned_markdown"] == source


def test_malformed_model_reply_retries_without_inserting_raw_text(tmp_path):
    image = tmp_path / "slide.png"
    image.write_bytes(b"image")
    reader = interpreter(iter([{}, reading(), review(issues=[{"reason": "不能核对"}])]))
    result = reader.read(image, "原始乱码", "", tmp_path / "cache")
    assert result["quality"] == "review_needed"
    assert result["raw_native"] == "原始乱码"


def test_visual_conversion_excludes_unreviewed_and_navigation_pages(tmp_path, monkeypatch):
    from io import BytesIO

    import pdfplumber
    from pptx import Presentation
    from pptx.util import Inches

    from final_review import material_conversion

    ppt = Presentation()
    for text in ["目录", "正常定义", "错字乱码"]:
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = text
    buffer = BytesIO()
    ppt.save(buffer)
    monkeypatch.setattr(
        material_conversion,
        "_convert_presentation_to_pdf",
        lambda _path, folder: folder / "test.pdf",
    )
    from contextlib import nullcontext

    monkeypatch.setattr(
        pdfplumber, "open", lambda _path: nullcontext(SimpleNamespace(pages=[None, None, None]))
    )
    monkeypatch.setattr(material_conversion, "_find_executable", lambda *_args: "pdftoppm")
    monkeypatch.setattr(
        material_conversion,
        "_run",
        lambda cmd, *_args: Path(cmd[-1]).with_suffix(".png").write_bytes(b"image"),
    )

    class Reader:
        def read(self, image, native, extracted, cache):
            return {
                "title": native,
                "kind": "navigation" if native == "目录" else "knowledge",
                "quality": "review_needed" if native == "错字乱码" else "verified",
                "issues": [],
                "blocks": [{"title": "整理后知识点", "markdown": "完整定义"}],
            }

    result = material_conversion.convert_material(
        buffer.getvalue(), "slides.pptx", 10**7, interpreter=Reader(), cache_dir=tmp_path / "pages"
    )
    assert len(result.sections) == 1
    assert result.sections[0]["position"] == 2
    assert "完整定义" in result.markdown
    assert "错字乱码" not in result.markdown
    assert "目录" not in result.markdown
    assert len(result.pages) == 3


def test_preview_exposes_quality_and_protects_original_page_path(system, tmp_path, monkeypatch):
    from final_review import material_jobs

    system.settings.uploads_dir = str(tmp_path)
    monkeypatch.setattr(
        material_jobs,
        "convert_material",
        lambda *_args, **_kwargs: ConvertedMaterial(
            markdown="整理后的真实知识",
            sections=[{"text": "整理后的真实知识", "position_kind": "slide", "position": 1}],
            pages=[
                {
                    "position": 1,
                    "title": "知识页",
                    "kind": "knowledge",
                    "quality": "verified",
                    "issues": [],
                }
            ],
            pipeline="visual-slides-v1",
        ),
    )
    with TestClient(create_app(system.settings, system)) as client:
        uploaded = client.post(
            "/knowledge/upload",
            data={"course_id": "net", "title": "课件", "source_type": "teacher_ppt"},
            files={"file": ("slides.pptx", b"test")},
        )
        job = system.store.claim_material_job()
        material_jobs.process_material_job(system.store, system.kb, job, 10**7)
        base = f"/api/courses/net/documents/{uploaded.json()['document_id']}"
        preview = client.get(base + "/chunks").json()
        assert preview["quality_status"] == "verified"
        document = system.store.get("document", uploaded.json()["document_id"])
        directory = Path(document["file_path"]).with_suffix(".analysis")
        directory.mkdir()
        (directory / "slide-001.png").write_bytes(b"PNG")
        assert client.get(base + "/pages/1").content == b"PNG"
        assert client.get(base + "/pages/2").status_code == 404
        assert client.get(base.replace("net", "other") + "/pages/1").status_code == 404


def test_billing_failure_is_actionable_and_never_indexes_raw_ppt(system, tmp_path, monkeypatch):
    from final_review import material_jobs
    from final_review.slide_understanding import VisionServiceUnavailable

    system.settings.uploads_dir = str(tmp_path)
    system.settings.material_vision_enabled = True

    def unavailable(_settings):
        raise VisionServiceUnavailable("视觉模型账号欠费，请恢复后重试")

    monkeypatch.setattr(material_jobs, "SlideInterpreter", unavailable)
    with TestClient(create_app(system.settings, system)) as client:
        response = client.post("/knowledge/upload", data={
            "course_id": "net", "title": "课件", "source_type": "teacher_ppt",
        }, files={"file": ("slides.pptx", b"raw ppt")})
        job = system.store.claim_material_job()
        material_jobs.process_material_job(system.store, system.kb, job, 10**7)
        result = system.store.get_material_job(job["job_id"])
        assert result["status"] == "failed"
        assert result["error_code"] == "vision_unavailable"
        assert "欠费" in result["error_message"]
        assert not any(chunk["document_id"] == response.json()["document_id"]
                       for chunk in system.store.chunks)
